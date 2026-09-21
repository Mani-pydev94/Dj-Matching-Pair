from io import BytesIO

from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import TestCase
from django.urls import reverse
from openpyxl import Workbook, load_workbook

from accounts.models import User
from admin_dashboard.views import _parse_questionnaire_workbook
from questionnaire.models import Question, QuestionOption, QuestionnaireCategory


class QuestionnaireAdminTests(TestCase):
    def setUp(self):
        self.user = User.objects.create_user(
            email='admin@example.com',
            password='Strong-password-123',
            is_staff=True,
        )
        self.category = QuestionnaireCategory.objects.create(
            name='Values',
            slug='values',
        )
        self.client.force_login(self.user)

    def test_question_edit_manages_options(self):
        response = self.client.post(
            reverse('admin_dashboard:question_new'),
            {
                'question_key': 'Q001',
                'category_ref': self.category.id,
                'text': 'What matters most?',
                'question_type': 'SINGLE_CHOICE',
                'help_text': '',
                'order': 1,
                'weight': '1.0',
                'is_required': 'on',
                'is_active': 'on',
                'options-TOTAL_FORMS': '2',
                'options-INITIAL_FORMS': '0',
                'options-MIN_NUM_FORMS': '0',
                'options-MAX_NUM_FORMS': '1000',
                'options-0-option_text': 'Communication',
                'options-0-option_value': 'communication',
                'options-0-compatibility_value': '5',
                'options-0-display_order': '1',
                'options-0-is_active': 'on',
                'options-1-option_text': 'Reliability',
                'options-1-option_value': 'reliability',
                'options-1-compatibility_value': '4',
                'options-1-display_order': '2',
                'options-1-is_active': 'on',
            },
        )

        self.assertEqual(response.status_code, 302)
        self.assertEqual(QuestionOption.objects.count(), 2)

    def test_export_contains_questionnaire_columns(self):
        response = self.client.get(reverse('admin_dashboard:export_questionnaire'))

        self.assertEqual(response.status_code, 200)
        workbook = load_workbook(BytesIO(response.content), read_only=True)
        self.assertEqual(workbook.active.cell(1, 1).value, 'question_key')

    def test_import_preview_does_not_write_until_confirmed(self):
        workbook = Workbook()
        sheet = workbook.active
        sheet.append([
            'question_key', 'category', 'question_text', 'question_type', 'help_text',
            'question_order', 'question_weight', 'is_required', 'is_active',
            'option_order', 'option_text', 'option_value', 'option_score',
            'option_is_active',
        ])
        sheet.append([
            'Q010', 'Values', 'What matters most?', 'SINGLE_CHOICE', '',
            1, 1.0, True, True, 1, 'Communication', 'communication', 5.0, True,
        ])
        file_data = BytesIO()
        workbook.save(file_data)
        rows, errors = _parse_questionnaire_workbook(
            SimpleUploadedFile(
                'questions.xlsx',
                file_data.getvalue(),
                content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
            ),
        )

        self.assertEqual(errors, [])
        self.assertEqual(rows[0]['question_key'], 'Q010')
        self.assertFalse(Question.objects.filter(question_key='Q010').exists())
