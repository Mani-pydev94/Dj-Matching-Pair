import csv
from io import BytesIO

from django.contrib import messages
from django.contrib.auth import get_user_model
from django.contrib.auth.models import Group, Permission
from django.core.paginator import Paginator
from django.db import transaction
from django.db.models import Count, Q
from django.http import HttpResponse
from django.shortcuts import get_object_or_404, redirect, render
from django.utils.text import slugify
from django.urls import reverse
from django.views.decorators.http import require_POST

from communities.models import Community
from connections.models import Connection
from events.models import Event
from notifications.models import Notification
from profiles.models import Profile
from questionnaire.models import Question, QuestionOption, QuestionResponse, QuestionnaireCategory

from .forms import (
    CategoryForm, CommunityForm, EventAdminForm, GroupForm, OptionFormSet,
    QuestionForm, UserAdminForm,
)
from .permissions import can_manage, dashboard_access


User = get_user_model()


def _role(user):
    if user.is_superuser:
        return 'SUPER ADMIN'
    return user.groups.values_list('name', flat=True).first() or ('ADMIN' if user.is_staff else 'STUDENT')


def _context(request, **extra):
    return {'admin_role': _role(request.user), **extra}


@dashboard_access
def dashboard(request):
    return render(request, 'admin_dashboard/dashboard.html', _context(
        request,
        stats={
            'users': User.objects.count(),
            'active_users': User.objects.filter(is_active=True).count(),
            'staff': User.objects.filter(is_staff=True).count(),
            'moderators': Group.objects.filter(name__iexact='Moderator').aggregate(
                total=Count('user')).get('total', 0),
            'communities': Community.objects.count(),
            'events': Event.objects.count(),
            'responses': QuestionResponse.objects.count(),
            'connections': Connection.objects.count(),
        },
        recent_users=User.objects.order_by('-created_at')[:6],
        recent_events=Event.objects.order_by('start')[:5],
    ))


@dashboard_access
def users(request):
    queryset = User.objects.select_related('profile').prefetch_related('groups').order_by('-created_at')
    term = request.GET.get('q', '').strip()
    if term:
        queryset = queryset.filter(
            Q(first_name__icontains=term) | Q(last_name__icontains=term) |
            Q(email__icontains=term) | Q(profile__university__icontains=term) |
            Q(profile__city__icontains=term)
        )
    status = request.GET.get('status')
    if status == 'active':
        queryset = queryset.filter(is_active=True)
    elif status == 'inactive':
        queryset = queryset.filter(is_active=False)
    page = Paginator(queryset, 20).get_page(request.GET.get('page'))
    return render(request, 'admin_dashboard/users.html', _context(request, page=page, query=term))


@dashboard_access
def user_detail(request, user_id):
    user = get_object_or_404(User.objects.select_related('profile').prefetch_related(
        'groups', 'communities', 'events'), id=user_id)
    return render(request, 'admin_dashboard/user_detail.html', _context(
        request, managed_user=user,
        responses=QuestionResponse.objects.filter(user=user).select_related('question'),
        connections=Connection.objects.filter(Q(requester=user) | Q(recipient=user)).select_related(
            'requester', 'recipient'),
    ))


@dashboard_access
def user_edit(request, user_id):
    if not can_manage(request, 'accounts.change_user'):
        from django.core.exceptions import PermissionDenied
        raise PermissionDenied
    user = get_object_or_404(User, id=user_id)
    form = UserAdminForm(request.POST or None, instance=user)
    if request.method == 'POST' and form.is_valid():
        form.save()
        messages.success(request, 'User updated.')
        return redirect('admin_dashboard:user_detail', user_id=user.id)
    return render(request, 'admin_dashboard/form.html', _context(request, form=form, title='Edit user', back_url=reverse('admin_dashboard:user_detail', args=[user.id])))


@dashboard_access
@require_POST
def toggle_user(request, user_id):
    if not can_manage(request, 'accounts.change_user'):
        from django.core.exceptions import PermissionDenied
        raise PermissionDenied
    user = get_object_or_404(User, id=user_id)
    if user == request.user:
        messages.error(request, 'You cannot deactivate your own account.')
    else:
        user.is_active = not user.is_active
        user.save(update_fields=('is_active',))
        messages.success(request, f'User {"activated" if user.is_active else "deactivated"}.')
    return redirect('admin_dashboard:user_detail', user_id=user.id)


@dashboard_access
def roles(request):
    return render(request, 'admin_dashboard/roles.html', _context(
        request, roles=Group.objects.annotate(user_count=Count('user')).prefetch_related('permissions'),
        permissions=Permission.objects.order_by('content_type__app_label', 'codename'),
    ))


@dashboard_access
def role_edit(request, role_id=None):
    if not request.user.is_superuser:
        from django.core.exceptions import PermissionDenied
        raise PermissionDenied
    role = get_object_or_404(Group, id=role_id) if role_id else None
    form = GroupForm(request.POST or None, instance=role)
    if request.method == 'POST' and form.is_valid():
        form.save()
        messages.success(request, 'Role saved.')
        return redirect('admin_dashboard:roles')
    return render(request, 'admin_dashboard/form.html', _context(
        request, form=form, title='Edit role' if role else 'Create role',
        back_url=reverse('admin_dashboard:roles'),
    ))


@dashboard_access
def events(request):
    return render(request, 'admin_dashboard/events.html', _context(request, events=Event.objects.all()))


@dashboard_access
def event_edit(request, event_id=None):
    event = get_object_or_404(Event, id=event_id) if event_id else None
    form = EventAdminForm(request.POST or None, instance=event)
    if request.method == 'POST' and form.is_valid():
        form.save()
        messages.success(request, 'Event saved.')
        return redirect('admin_dashboard:events')
    return render(request, 'admin_dashboard/form.html', _context(request, form=form, title='Edit event' if event else 'Create event', back_url=reverse('admin_dashboard:events')))


@dashboard_access
@require_POST
def event_delete(request, event_id):
    if not request.user.is_superuser:
        from django.core.exceptions import PermissionDenied
        raise PermissionDenied
    get_object_or_404(Event, id=event_id).delete()
    messages.success(request, 'Event deleted.')
    return redirect('admin_dashboard:events')


@dashboard_access
def communities(request):
    return render(request, 'admin_dashboard/communities.html', _context(request, communities=Community.objects.all()))


@dashboard_access
def community_edit(request, community_id=None):
    community = get_object_or_404(Community, id=community_id) if community_id else None
    form = CommunityForm(request.POST or None, instance=community)
    if request.method == 'POST' and form.is_valid():
        form.save()
        messages.success(request, 'Community saved.')
        return redirect('admin_dashboard:communities')
    return render(request, 'admin_dashboard/form.html', _context(request, form=form, title='Edit community' if community else 'Create community', back_url=reverse('admin_dashboard:communities')))


@dashboard_access
def questionnaire(request):
    return render(request, 'admin_dashboard/questionnaire.html', _context(
        request, categories=QuestionnaireCategory.objects.annotate(question_count=Count('questions')),
        questions=Question.objects.select_related('category_ref').order_by('category_ref__display_order', 'order'),
        responses=QuestionResponse.objects.count(),
    ))


@dashboard_access
def questionnaire_category_edit(request, category_id=None):
    category = get_object_or_404(QuestionnaireCategory, id=category_id) if category_id else None
    form = CategoryForm(request.POST or None, instance=category)
    if request.method == 'POST' and form.is_valid():
        form.save()
        messages.success(request, 'Questionnaire section saved.')
        return redirect('admin_dashboard:questionnaire')
    return render(request, 'admin_dashboard/form.html', _context(request, form=form, title='Edit section' if category else 'Create section', back_url=reverse('admin_dashboard:questionnaire')))


@dashboard_access
def questionnaire_question_edit(request, question_id=None):
    question = get_object_or_404(Question, id=question_id) if question_id else None
    form = QuestionForm(request.POST or None, instance=question)
    option_formset = OptionFormSet(
        request.POST or None,
        instance=question or Question(),
        prefix='options',
    )
    if request.method == 'POST' and form.is_valid():
        with transaction.atomic():
            saved_question = form.save(commit=False)
            if not saved_question.question_key:
                next_number = Question.objects.count() + 1
                while Question.objects.filter(question_key=f'Q{next_number:03d}').exists():
                    next_number += 1
                saved_question.question_key = f'Q{next_number:03d}'
            saved_question.save()
            option_formset.instance = saved_question
            if option_formset.is_valid():
                protected_options = []
                for deleted_form in option_formset.deleted_forms:
                    if deleted_form.instance.pk and deleted_form.instance.responses.exists():
                        deleted_form.instance.is_active = False
                        deleted_form.instance.save(update_fields=('is_active',))
                        deleted_form.cleaned_data['DELETE'] = False
                        protected_options.append(deleted_form.instance.option_value)
                option_formset.save()
                if protected_options:
                    messages.warning(
                        request,
                        'Options with existing responses were deactivated instead of deleted: '
                        + ', '.join(protected_options),
                    )
                messages.success(request, 'Question and answer options saved.')
                return redirect('admin_dashboard:questionnaire')
    return render(request, 'admin_dashboard/question_edit.html', _context(
        request,
        form=form,
        option_formset=option_formset,
        question=question,
        title='Edit question' if question else 'Create question',
        back_url=reverse('admin_dashboard:questionnaire'),
    ))


QUESTIONNAIRE_COLUMNS = (
    'question_key', 'category', 'question_text', 'question_type', 'help_text',
    'question_order', 'question_weight', 'is_required', 'is_active',
    'option_order', 'option_text', 'option_value', 'option_score',
    'option_is_active',
)


def _excel_bool(value, field_name, row_number, errors):
    if isinstance(value, bool):
        return value
    normalized = str(value).strip().lower()
    if normalized in ('true', '1', 'yes'):
        return True
    if normalized in ('false', '0', 'no'):
        return False
    errors.append(f'Row {row_number}: {field_name} must be TRUE or FALSE.')
    return False


def _parse_questionnaire_workbook(upload):
    from openpyxl import load_workbook

    if not upload:
        return [], ['Choose an .xlsx workbook to import.']
    if not upload.name.lower().endswith('.xlsx'):
        return [], ['Only .xlsx files are supported.']
    try:
        workbook = load_workbook(upload, read_only=True, data_only=True)
    except Exception as exc:
        return [], [f'Unable to read Excel file: {exc}']
    sheet = workbook.active
    rows = list(sheet.iter_rows(values_only=True))
    if not rows:
        return [], ['The workbook is empty.']
    headers = [str(value).strip() if value is not None else '' for value in rows[0]]
    missing = [column for column in QUESTIONNAIRE_COLUMNS if column not in headers]
    if missing:
        return [], [f'Missing required columns: {", ".join(missing)}']
    indexes = {column: headers.index(column) for column in QUESTIONNAIRE_COLUMNS}
    parsed, errors = [], []
    valid_types = {'SINGLE_CHOICE', 'MULTIPLE_CHOICE', 'SCALE', 'YES_NO'}
    for row_number, row in enumerate(rows[1:], start=2):
        if not any(value not in (None, '') for value in row):
            continue
        data = {column: row[indexes[column]] for column in QUESTIONNAIRE_COLUMNS}
        for field in ('question_key', 'category', 'question_text', 'question_type'):
            if not str(data[field] or '').strip():
                errors.append(f'Row {row_number}: Missing {field}.')
        data['question_type'] = str(data['question_type'] or '').strip().upper()
        if data['question_type'] not in valid_types:
            errors.append(f'Row {row_number}: Invalid question_type.')
        if not str(data['option_text'] or '').strip() or not str(data['option_value'] or '').strip():
            errors.append(f'Row {row_number}: option_text and option_value are required.')
        try:
            data['question_order'] = int(data['question_order'])
            data['option_order'] = int(data['option_order'])
            data['question_weight'] = float(data['question_weight'])
            data['option_score'] = float(data['option_score'])
        except (TypeError, ValueError):
            errors.append(f'Row {row_number}: order fields must be integers and score/weight must be numeric.')
        data['is_required'] = _excel_bool(data['is_required'], 'is_required', row_number, errors)
        data['is_active'] = _excel_bool(data['is_active'], 'is_active', row_number, errors)
        data['option_is_active'] = _excel_bool(data['option_is_active'], 'option_is_active', row_number, errors)
        data['_row_number'] = row_number
        parsed.append(data)
    grouped = {}
    for data in parsed:
        key = str(data['question_key']).strip()
        grouped.setdefault(key, []).append(data)
    for key, items in grouped.items():
        values = [str(item['option_value']).strip() for item in items]
        if len(values) != len(set(values)):
            errors.append(f'Duplicate option_value for {key}.')
    return parsed, errors


@dashboard_access
def questionnaire_import_export(request):
    preview = None
    errors = []
    if request.method == 'POST' and request.POST.get('action') == 'preview':
        parsed, errors = _parse_questionnaire_workbook(request.FILES.get('workbook'))
        if not errors:
            request.session['questionnaire_import_rows'] = parsed
            keys = {item['question_key'] for item in parsed}
            existing = set(Question.objects.filter(question_key__in=keys).values_list('question_key', flat=True))
            preview = {
                'questions': len(keys),
                'new_questions': len(keys - existing),
                'updates': len(keys & existing),
                'options': len(parsed),
                'rows': parsed,
            }
    elif request.method == 'POST' and request.POST.get('action') == 'import':
        rows = request.session.pop('questionnaire_import_rows', [])
        mode = request.POST.get('mode', 'update')
        if not rows:
            errors = ['Upload and preview a workbook before importing.']
        elif mode == 'create' and Question.objects.filter(
            question_key__in={row['question_key'] for row in rows},
        ).exists():
            errors = ['Create Only cannot import an existing question_key.']
        else:
            with transaction.atomic():
                for key, items in _group_import_rows(rows).items():
                    first = items[0]
                    category, _ = QuestionnaireCategory.objects.get_or_create(
                        name=str(first['category']).strip(),
                        defaults={'slug': slugify(str(first['category']))},
                    )
                    question, created = Question.objects.update_or_create(
                        question_key=key,
                        defaults={
                            'category': category.name,
                            'category_ref': category,
                            'text': first['question_text'],
                            'question_type': first['question_type'],
                            'help_text': first['help_text'] or '',
                            'order': first['question_order'],
                            'weight': first['question_weight'],
                            'is_required': first['is_required'],
                            'is_active': first['is_active'],
                        },
                    )
                    for item in items:
                        QuestionOption.objects.update_or_create(
                            question=question,
                            option_value=str(item['option_value']).strip(),
                            defaults={
                                'option_text': item['option_text'],
                                'display_order': item['option_order'],
                                'compatibility_value': item['option_score'],
                                'is_active': item['option_is_active'],
                            },
                        )
            messages.success(request, f'Imported {len(rows)} answer-option rows.')
            return redirect('admin_dashboard:questionnaire_import_export')
    return render(request, 'admin_dashboard/questionnaire_import_export.html', _context(
        request, preview=preview, errors=errors,
        categories=QuestionnaireCategory.objects.order_by('display_order', 'name'),
    ))


def _group_import_rows(rows):
    grouped = {}
    for row in rows:
        grouped.setdefault(str(row['question_key']).strip(), []).append(row)
    return grouped


@dashboard_access
def export_questionnaire(request):
    from openpyxl import Workbook

    questions = Question.objects.select_related('category_ref').prefetch_related('options').order_by('category', 'order')
    category_id = request.GET.get('category')
    question_type = request.GET.get('question_type')
    status = request.GET.get('status')
    required = request.GET.get('required')
    if category_id:
        questions = questions.filter(category_ref_id=category_id)
    if question_type:
        questions = questions.filter(question_type=question_type)
    if status in ('active', 'inactive'):
        questions = questions.filter(is_active=status == 'active')
    if required in ('required', 'optional'):
        questions = questions.filter(is_required=required == 'required')
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = 'Questionnaire'
    sheet.append(QUESTIONNAIRE_COLUMNS)
    if request.GET.get('template'):
        example_rows = [
            ['Q001', 'Values', 'What matters most when working with a teammate?', 'SINGLE_CHOICE', 'Choose one', 1, 1.0, True, True, 1, 'Clear communication', 'communication', 5.0, True],
            ['Q001', 'Values', 'What matters most when working with a teammate?', 'SINGLE_CHOICE', 'Choose one', 1, 1.0, True, True, 2, 'Technical expertise', 'technical_expertise', 4.0, True],
        ]
        for row in example_rows:
            sheet.append(row)
    else:
        for question in questions:
            options = list(question.options.all()) or [None]
            for option in options:
                sheet.append([
                    question.question_key or f'Q{question.id}',
                    question.category_ref.name if question.category_ref else question.category,
                    question.text, question.question_type, question.help_text,
                    question.order, question.weight, question.is_required, question.is_active,
                    option.display_order if option else 0,
                    option.option_text if option else '', option.option_value if option else '',
                    option.compatibility_value if option else 0,
                    option.is_active if option else True,
                ])
    instructions = workbook.create_sheet('Instructions')
    instructions.append(['Questionnaire import instructions'])
    for line in (
        'One row represents one answer option.',
        'Use the same question_key for every option belonging to one question.',
        'option_value must be unique within a question.',
        'Use TRUE or FALSE for boolean fields.',
        'score and weight must be numeric.',
        'Supported question types: SINGLE_CHOICE, MULTIPLE_CHOICE, SCALE, YES_NO.',
    ):
        instructions.append([line])
    output = BytesIO()
    workbook.save(output)
    response = HttpResponse(
        output.getvalue(),
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
    )
    response['Content-Disposition'] = 'attachment; filename="questionnaire.xlsx"'
    return response


@dashboard_access
def matches(request):
    return render(request, 'admin_dashboard/matches.html', _context(
        request, matches=Connection.objects.select_related('requester', 'recipient').order_by('-created_at'),
    ))


@dashboard_access
def notifications(request):
    return render(request, 'admin_dashboard/notifications.html', _context(
        request, notifications=Notification.objects.select_related('user')[:100],
    ))


@dashboard_access
def event_registrations(request):
    registrations = [
        (event, attendee)
        for event in Event.objects.prefetch_related('attendees')
        for attendee in event.attendees.all()
    ]
    return render(request, 'admin_dashboard/event_registrations.html', _context(
        request, registrations=registrations,
    ))


@dashboard_access
def community_joins(request):
    memberships = [
        (community, member)
        for community in Community.objects.prefetch_related('members')
        for member in community.members.all()
    ]
    return render(request, 'admin_dashboard/community_joins.html', _context(
        request, memberships=memberships,
    ))


@dashboard_access
def reports(request):
    return render(request, 'admin_dashboard/reports.html', _context(request))


@dashboard_access
def export_users(request):
    response = HttpResponse(content_type='text/csv')
    response['Content-Disposition'] = 'attachment; filename="campus-users.csv"'
    writer = csv.writer(response)
    writer.writerow(('Name', 'Email', 'University', 'City', 'Active', 'Joined'))
    for user in User.objects.select_related('profile'):
        profile = getattr(user, 'profile', None)
        writer.writerow((user.get_full_name(), user.email, profile.university if profile else '', profile.city if profile else '', user.is_active, user.created_at.isoformat()))
    return response
